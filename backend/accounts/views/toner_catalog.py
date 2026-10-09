import json, logging, os, re, threading
import openpyxl
from django.conf import settings
from django.http import JsonResponse
from django.views.decorators.csrf import csrf_exempt

logger = logging.getLogger(__name__)

# Kept under backend/data/ (not STATIC_ROOT or MEDIA_ROOT, so it's never
# served publicly) and overridable for deployments that move the file.
DEFAULT_XLSX_PATH = os.path.join(settings.BASE_DIR, "data", "HP_Toner_Printer_Compatibility.xlsx")

_lock = threading.Lock()
_cache = None  # {"toners_by_brand": {...}, "printers_by_key": {...}} once loaded
_cache_signature = None


def _xlsx_path():
    return os.environ.get("TONER_COMPAT_XLSX_PATH", DEFAULT_XLSX_PATH)


def _sheet_rows(wb, sheet_name):
    ws = wb[sheet_name]
    rows = ws.iter_rows(values_only=True)
    header = next(rows)
    idx = {name: i for i, name in enumerate(header) if name}
    for row in rows:
        if row and any(cell not in (None, "") for cell in row):
            yield {name: row[i] if i < len(row) else None for name, i in idx.items()}


def _load():
    path = _xlsx_path()
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if "Toner_Master" not in wb.sheetnames:
            return _load_combined_sheet(wb)
        toner_order = []
        cartridge_type = {}
        part_numbers = {}
        for row in _sheet_rows(wb, "Toner_Master"):
            brand = str(row.get("Brand") or "").strip()
            model = str(row.get("Toner Model") or "").strip()
            if not brand or not model:
                continue
            key = (brand, model)
            if key not in cartridge_type:
                cartridge_type[key] = str(row.get("Cartridge Type") or "").strip()
                toner_order.append(key)
            color = str(row.get("Cartridge Color") or "").strip()
            part_no = str(row.get("Part No.") or "").strip()
            entries = part_numbers.setdefault(key, [])
            if not any(e["color"] == color and e["partNo"] == part_no for e in entries):
                entries.append({"color": color, "partNo": part_no})

        toners_by_brand = {}
        for brand, model in toner_order:
            toners_by_brand.setdefault(brand, []).append({
                "brand": brand,
                "tonerModel": model,
                "cartridgeType": cartridge_type[(brand, model)],
                "partNumbers": part_numbers.get((brand, model), []),
            })

        printers_by_key = {}
        seen_printer = {}
        for row in _sheet_rows(wb, "Compatibility"):
            brand = str(row.get("Brand") or "").strip()
            toner_model = str(row.get("Toner Model") or "").strip()
            printer_model = str(row.get("Printer Model") or "").strip()
            printer_type = str(row.get("Printer Type") or "").strip()
            if not brand or not toner_model or not printer_model:
                continue
            key = (brand, toner_model)
            seen = seen_printer.setdefault(key, set())
            if printer_model in seen:
                continue
            seen.add(printer_model)
            printers_by_key.setdefault(key, []).append({
                "brand": brand, "printerModel": printer_model, "printerType": printer_type,
            })

        for key in printers_by_key:
            printers_by_key[key].sort(key=lambda p: p["printerModel"])

        return {"toners_by_brand": toners_by_brand, "printers_by_key": printers_by_key}
    finally:
        wb.close()


def _load_combined_sheet(wb):
    toners_by_brand, printers_by_key = {}, {}
    for sheet in wb.sheetnames:
        rows = wb[sheet].iter_rows(values_only=True)
        header = next(rows, ())
        columns = {str(value).strip().casefold(): i for i, value in enumerate(header) if value}
        required = ("toner model", "printer brand", "compatible printer models")
        if not all(name in columns for name in required):
            continue
        for row in rows:
            def cell(name):
                index = columns[name]
                return str(row[index] or "").strip() if index < len(row) else ""
            model = cell("toner model")
            brands = [brand.strip() for brand in cell("printer brand").split("/") if brand.strip()]
            if not model or not brands:
                continue
            printer_models = list(dict.fromkeys(
                value.strip() for value in re.split(r"[,;\n/]+", cell("compatible printer models")) if value.strip()
            ))
            for brand in brands:
                key = (brand, model)
                if key not in printers_by_key:
                    toners_by_brand.setdefault(brand, []).append({
                        "brand": brand, "tonerModel": model, "cartridgeType": "",
                        "partNumbers": [{"color": "", "partNo": model}],
                    })
                    printers_by_key[key] = []
                seen = {printer["printerModel"] for printer in printers_by_key[key]}
                printers_by_key[key].extend(
                    {"brand": brand, "printerModel": printer, "printerType": ""}
                    for printer in printer_models if printer not in seen
                )
        for printers in printers_by_key.values():
            printers.sort(key=lambda printer: printer["printerModel"])
        return {"toners_by_brand": toners_by_brand, "printers_by_key": printers_by_key}
    raise KeyError("Expected toner model, Printer Brand and Compatible Printer Models columns")


def _get_cache():
    global _cache, _cache_signature
    path = _xlsx_path()
    stat = os.stat(path)
    signature = (path, stat.st_mtime_ns, stat.st_size)
    with _lock:
        if _cache is None or _cache_signature != signature:
            _cache = _load()
            _cache_signature = signature
        return _cache


def get_toners_for_brands(brands):
    cache = _get_cache()
    out = []
    for brand in brands:
        out.extend(cache["toners_by_brand"].get(brand, []))
    return out


def get_compatible_printers(brand, toner_model):
    cache = _get_cache()
    return cache["printers_by_key"].get((brand, toner_model), [])


def get_printers_for_toners(toners):
    """toners: list of {"brand":..., "tonerModel":...}. Returns (per-toner results, de-duplicated union)."""
    results = []
    union = {}
    for item in toners:
        brand = str(item.get("brand") or "").strip()
        toner_model = str(item.get("tonerModel") or "").strip()
        printers = get_compatible_printers(brand, toner_model)
        results.append({"brand": brand, "tonerModel": toner_model, "printers": printers})
        for p in printers:
            union.setdefault(p["printerModel"], p)
    union_list = sorted(union.values(), key=lambda p: p["printerModel"])
    return results, union_list


def toner_catalog_toners(r):
    brands_param = (r.GET.get("brands") or r.GET.get("brand") or "").strip()
    brands = [b.strip() for b in brands_param.split(",") if b.strip()]
    if not brands:
        return JsonResponse({"error": "brands query param is required"}, status=400)
    try:
        toners = get_toners_for_brands(brands)
    except FileNotFoundError:
        logger.error("Toner compatibility workbook not found at %s", _xlsx_path())
        return JsonResponse({"error": "Toner/printer compatibility data is unavailable"}, status=503)
    except KeyError as e:
        logger.error("Toner compatibility workbook missing expected sheet/column: %s", e)
        return JsonResponse({"error": "Toner/printer compatibility data is malformed"}, status=503)
    except Exception:
        logger.exception("Failed to read toner compatibility workbook")
        return JsonResponse({"error": "Toner/printer compatibility data is unavailable"}, status=503)
    return JsonResponse({"toners": toners, "brands": list(_get_cache()["toners_by_brand"])})


def toner_catalog_printers(r):
    brand = (r.GET.get("brand") or "").strip()
    toner_model = (r.GET.get("tonerModel") or "").strip()
    if not brand or not toner_model:
        return JsonResponse({"error": "brand and tonerModel query params are required"}, status=400)
    try:
        printers = get_compatible_printers(brand, toner_model)
    except FileNotFoundError:
        logger.error("Toner compatibility workbook not found at %s", _xlsx_path())
        return JsonResponse({"error": "Toner/printer compatibility data is unavailable"}, status=503)
    except KeyError as e:
        logger.error("Toner compatibility workbook missing expected sheet/column: %s", e)
        return JsonResponse({"error": "Toner/printer compatibility data is malformed"}, status=503)
    except Exception:
        logger.exception("Failed to read toner compatibility workbook")
        return JsonResponse({"error": "Toner/printer compatibility data is unavailable"}, status=503)
    return JsonResponse({"printers": printers})


@csrf_exempt
def toner_catalog_printers_bulk(r):
    if r.method != "POST":
        return JsonResponse({"error": "POST required"}, status=405)
    try:
        payload = json.loads(r.body or "{}")
    except (json.JSONDecodeError, TypeError):
        return JsonResponse({"error": "Invalid JSON body"}, status=400)
    toners = payload.get("toners")
    if not isinstance(toners, list) or not toners:
        return JsonResponse({"error": "toners (a non-empty list of {brand, tonerModel}) is required"}, status=400)
    for item in toners:
        if not isinstance(item, dict) or not str(item.get("brand") or "").strip() or not str(item.get("tonerModel") or "").strip():
            return JsonResponse({"error": "Each toner must have a non-empty brand and tonerModel"}, status=400)
    try:
        results, union_list = get_printers_for_toners(toners)
    except FileNotFoundError:
        logger.error("Toner compatibility workbook not found at %s", _xlsx_path())
        return JsonResponse({"error": "Toner/printer compatibility data is unavailable"}, status=503)
    except KeyError as e:
        logger.error("Toner compatibility workbook missing expected sheet/column: %s", e)
        return JsonResponse({"error": "Toner/printer compatibility data is malformed"}, status=503)
    except Exception:
        logger.exception("Failed to read toner compatibility workbook")
        return JsonResponse({"error": "Toner/printer compatibility data is unavailable"}, status=503)
    return JsonResponse({"results": results, "printers": union_list})
