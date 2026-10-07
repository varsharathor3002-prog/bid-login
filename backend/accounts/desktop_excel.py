"""Two-way sync between the Desktop product directory and Desktop_Product.xlsx.

* Every Desktop CatalogueProduct save/delete (UI, bid flow, GeM extension) is
  written to the Excel sheet through model signals (see apps.py).
* When the directory is read, a changed Excel file is imported: new rows are
  created and rows edited in Excel (file newer than the product) are updated.
  Desktop products missing from the sheet are appended to it.
* The very first sync on a server (no Desktop_Product.xlsx.synced marker)
  backs the sheet up and rewrites it from the directory, so an older sheet
  never overwrites newer directory edits.

Two sheet layouts are read and written: the old DB export (specs as JSON in
"extra_specs") and the readable one written by replace_directory (one column
per spec field, e.g. "Type of RAM"). Spec columns win over the JSON.
"""
import json
import logging
import os
import shutil
import sys
import tempfile
import threading
import time
from contextlib import contextmanager
from datetime import datetime

import openpyxl
from django.conf import settings
from django.db import transaction

logger = logging.getLogger(__name__)

TEXT_FIELDS = ["category", "description", "os", "processor", "ram", "storage"]
DEFAULT_HEADERS = [
    "id", "model_no", "category", "image", "created_at", "updated_at",
    "extra_specs", "description", "os", "processor", "ram", "storage",
    "gem_product_id",
]
LOCK_STALE_SECONDS = 60

_local = threading.local()
_process_lock = threading.Lock()
_last_seen_mtime = None
# The test runner must never touch the real sheet; tests patch enabled().
_RUNNING_TESTS = len(sys.argv) > 1 and sys.argv[1] == "test"


def enabled():
    return not _RUNNING_TESTS


def excel_path():
    return os.path.join(settings.BASE_DIR, "Desktop_Product.xlsx")


def is_desktop(product):
    return (product.category or "").strip().lower() in {"", "desktop"}


@contextmanager
def paused():
    """Skip the per-save signal writes (bulk operations write the sheet once)."""
    previous = getattr(_local, "paused", False)
    _local.paused = True
    try:
        yield
    finally:
        _local.paused = previous


def is_paused():
    return getattr(_local, "paused", False)


@contextmanager
def _file_lock(path):
    # Gunicorn workers are separate processes: a lock file serialises their
    # read-modify-write of the sheet so no worker overwrites another's row.
    lock_path = f"{path}.lock"
    deadline = time.time() + 30
    with _process_lock:
        while True:
            try:
                fd = os.open(lock_path, os.O_CREAT | os.O_EXCL | os.O_WRONLY)
                os.close(fd)
                break
            except FileExistsError:
                try:
                    if time.time() - os.path.getmtime(lock_path) > LOCK_STALE_SECONDS:
                        os.remove(lock_path)
                        continue
                except OSError:
                    continue
                if time.time() > deadline:
                    raise TimeoutError("Desktop_Product.xlsx is busy. Try again.")
                time.sleep(0.1)
        try:
            yield
        finally:
            try:
                os.remove(lock_path)
            except OSError:
                pass


def _cell_text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).strip()
    return "" if text.lower() in {"nan", "none", "null"} else text


def _spec_fields():
    from .views.Desktop import ALL_FIELDS
    return ALL_FIELDS


def _product_specs(product):
    specs = product.extra_specs or {}
    if isinstance(specs, str):
        try:
            specs = json.loads(specs)
        except ValueError:
            specs = {}
    return specs if isinstance(specs, dict) else {}


def product_key(model_no, gem_product_id=""):
    """A product is its model number plus GeM listing id: GeM lists several
    configurations under one model number."""
    return ((model_no or "").strip().upper(), (gem_product_id or "").strip())


def _key_of(product):
    return product_key(product.model_no, product.gem_product_id)


def _match_products(keys, products):
    """{key: existing product} for GeM keys. A product saved before GeM ids
    were kept (same model, no id) is taken over by that model's first listing."""
    matched, used = {}, set()
    for key in keys:
        product = products.get(key)
        if product is None and key[1]:
            legacy = products.get((key[0], ""))
            if legacy is not None and legacy.pk not in used:
                product = legacy
        if product is not None and product.pk not in used:
            used.add(product.pk)
            matched[key] = product
    return matched


def _row_values(product):
    specs = _product_specs(product)
    values = {
        "id": str(product.id),
        "model_no": product.model_no or "",
        "gem_product_id": product.gem_product_id or "",
        "image": product.image.name if product.image else None,
        "created_at": str(product.created_at) if product.created_at else "",
        "updated_at": str(product.updated_at) if product.updated_at else "",
        "extra_specs": json.dumps(specs),
        **{field: getattr(product, field) or "" for field in TEXT_FIELDS},
    }
    for field in _spec_fields():
        values[field] = specs.get(field, "")
    return values


def _headers(sheet):
    return [_cell_text(cell.value) for cell in sheet[1]]


def _key_rows(sheet, headers):
    """{(MODEL_NO, gem id): [row numbers]} for every data row with a model number."""
    col = headers.index("model_no") + 1
    gem_col = headers.index("gem_product_id") + 1 if "gem_product_id" in headers else None
    rows = {}
    for row in range(2, sheet.max_row + 1):
        model = _cell_text(sheet.cell(row, col).value)
        if model:
            gem_id = _cell_text(sheet.cell(row, gem_col).value) if gem_col else ""
            rows.setdefault(product_key(model, gem_id), []).append(row)
    return rows


def _write_row(sheet, headers, row, values):
    for col, header in enumerate(headers, start=1):
        if header in values:
            sheet.cell(row, col).value = values[header]


def _save_workbook(workbook, path):
    # Write a temp file first so a failed save never corrupts the sheet.
    fd, temp_path = tempfile.mkstemp(suffix=".xlsx", dir=os.path.dirname(path))
    os.close(fd)
    try:
        workbook.save(temp_path)
        os.replace(temp_path, path)
    finally:
        if os.path.exists(temp_path):
            os.remove(temp_path)


def _open(path):
    if os.path.exists(path):
        workbook = openpyxl.load_workbook(path)
        sheet = workbook.active
        headers = _headers(sheet)
        if "model_no" not in headers:
            raise ValueError("Desktop_Product.xlsx has no model_no column.")
        if "gem_product_id" not in headers:
            headers.append("gem_product_id")
            sheet.cell(1, len(headers)).value = "gem_product_id"
        return workbook, sheet, headers
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Desktop Products"
    sheet.append(DEFAULT_HEADERS)
    return workbook, sheet, list(DEFAULT_HEADERS)


def _remember_mtime(path):
    global _last_seen_mtime
    try:
        _last_seen_mtime = os.path.getmtime(path)
    except OSError:
        _last_seen_mtime = None


def _edit_sheet(change):
    """Run change(sheet, headers) under the file lock and save. Returns error text or ""."""
    path = excel_path()
    try:
        with _file_lock(path):
            mtime_before = os.path.getmtime(path) if os.path.exists(path) else None
            workbook, sheet, headers = _open(path)
            if change(sheet, headers) is False:
                return ""
            _save_workbook(workbook, path)
            # If the sheet had unseen hand edits, leave them for the next import.
            if mtime_before is None or mtime_before == _last_seen_mtime:
                _remember_mtime(path)
        return ""
    except PermissionError:
        return "Desktop_Product.xlsx is open in another program. Close it and save again."
    except Exception as e:
        logger.exception("Desktop_Product.xlsx update failed")
        return f"Desktop_Product.xlsx could not be updated: {e}"


def save_product(product):
    """Write one Desktop product's row (update in place, or append). Returns error text or ""."""
    if not enabled() or not is_desktop(product):
        return ""
    values = _row_values(product)

    def change(sheet, headers):
        rows = _key_rows(sheet, headers).get(_key_of(product))
        for row in rows or [sheet.max_row + 1]:
            _write_row(sheet, headers, row, values)

    return _edit_sheet(change)


def remove_products(keys):
    """Delete the rows of these (model number, gem id) keys. Returns error text or ""."""
    targets = {product_key(*key) for key in keys if key and key[0]}
    if not enabled() or not targets:
        return ""

    def change(sheet, headers):
        rows = sorted(
            (row for key, numbers in _key_rows(sheet, headers).items() if key in targets for row in numbers),
            reverse=True,
        )
        if not rows:
            return False
        for row in rows:
            sheet.delete_rows(row)

    return _edit_sheet(change)


def backup_and_clear():
    """Used by Delete All: keep a dated backup, then empty the sheet's data rows."""
    path = excel_path()
    if not enabled() or not os.path.exists(path):
        return ""
    _backup(path)

    def change(sheet, headers):
        if sheet.max_row > 1:
            sheet.delete_rows(2, sheet.max_row - 1)

    return _edit_sheet(change)


def _row_product_data(headers, row):
    data = {header: _cell_text(value) for header, value in zip(headers, row) if header}
    specs = {}
    if data.get("extra_specs"):
        try:
            parsed = json.loads(data["extra_specs"])
            if isinstance(parsed, dict):
                specs = parsed
        except ValueError:
            pass
    # Spec columns win over the JSON; an emptied cell clears that spec.
    for field in _spec_fields():
        if field in data:
            specs[field] = data[field]
    return data, specs


def _filled(specs):
    return {k: v for k, v in specs.items() if v not in ("", None)}


def _baseline_marker():
    return f"{excel_path()}.synced"


def _backup(path):
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    shutil.copy2(path, os.path.join(os.path.dirname(path), f"Desktop_Product.backup-{stamp}.xlsx"))


def _baseline():
    """First sync on a server: the directory (DB) is what users see, so rewrite
    every Desktop row of the sheet from it instead of letting an older sheet
    overwrite newer directory edits. Excel-only rows are still imported."""
    from .models import CatalogueProduct

    path = excel_path()
    with _file_lock(path):
        _backup(path)
        workbook, sheet, headers = _open(path)
        rows = _key_rows(sheet, headers)
        products = {_key_of(p): p for p in CatalogueProduct.objects.all()}
        created = 0
        with paused():
            for key, numbers in rows.items():
                if key in products:
                    continue
                data, specs = _row_product_data(headers, [c.value for c in sheet[numbers[0]]])
                products[key] = CatalogueProduct.objects.create(
                    model_no=data["model_no"], gem_product_id=key[1], category=data.get("category") or "Desktop",
                    extra_specs=specs, image=data.get("image") or None,
                    **{f: data.get(f, "") for f in TEXT_FIELDS if f != "category"},
                )
                created += 1
        written = 0
        for key, product in products.items():
            if not is_desktop(product):
                continue
            for row in rows.get(key) or [sheet.max_row + 1]:
                _write_row(sheet, headers, row, _row_values(product))
            written += 1
        _save_workbook(workbook, path)
        with open(_baseline_marker(), "w", encoding="utf-8") as marker:
            marker.write(datetime.now().isoformat())
        _remember_mtime(path)
    return {"created": created, "updated": 0, "appended": written, "baseline": True}


def sync_from_excel(force=False):
    """Import Excel changes into the directory when the sheet changed since last seen."""
    from .models import CatalogueProduct

    path = excel_path()
    if not enabled() or not os.path.exists(path):
        return {"created": 0, "updated": 0, "appended": 0}
    if not os.path.exists(_baseline_marker()):
        try:
            return _baseline()
        except Exception:
            logger.exception("Desktop_Product.xlsx first sync failed")
            return {"created": 0, "updated": 0, "appended": 0, "error": True}
    mtime = os.path.getmtime(path)
    if not force and mtime == _last_seen_mtime:
        return {"created": 0, "updated": 0, "appended": 0}

    created = updated = 0
    try:
        with _file_lock(path):
            mtime = os.path.getmtime(path)
            workbook, sheet, headers = _open(path)
            products = {_key_of(p): p for p in CatalogueProduct.objects.all()}
            seen = set()
            with paused():
                for row in sheet.iter_rows(min_row=2, values_only=True):
                    data, specs = _row_product_data(headers, row)
                    model_no = data.get("model_no", "").strip()
                    key = product_key(model_no, data.get("gem_product_id"))
                    if not key[0] or key in seen:
                        continue
                    seen.add(key)
                    product = products.get(key)
                    if product is None:
                        with transaction.atomic():
                            products[key] = CatalogueProduct.objects.create(
                                model_no=model_no,
                                gem_product_id=key[1],
                                category=data.get("category") or "Desktop",
                                extra_specs=specs,
                                image=data.get("image") or None,
                                **{f: data.get(f, "") for f in TEXT_FIELDS if f != "category"},
                            )
                        created += 1
                        continue
                    # Only Excel edits newer than the product win; a product
                    # saved after the sheet was written keeps its own values.
                    if not is_desktop(product) or product.updated_at.timestamp() >= mtime:
                        continue
                    changed = []
                    for field in TEXT_FIELDS:
                        if field in data and data[field] != (getattr(product, field) or ""):
                            setattr(product, field, data[field])
                            changed.append(field)
                    current = _product_specs(product)
                    if _filled({**current, **specs}) != _filled(current):
                        product.extra_specs = {**current, **specs}
                        changed.append("extra_specs")
                    if data.get("image") and data["image"] != (product.image.name if product.image else ""):
                        product.image = data["image"]
                        changed.append("image")
                    if changed:
                        product.save(update_fields=[*changed, "updated_at"])
                        updated += 1

            missing = [p for key, p in products.items() if key not in seen and is_desktop(p)]
            for product in missing:
                _write_row(sheet, headers, sheet.max_row + 1, _row_values(product))
            if missing:
                _save_workbook(workbook, path)
            _remember_mtime(path)
        return {"created": created, "updated": updated, "appended": len(missing)}
    except Exception:
        logger.exception("Desktop_Product.xlsx import failed")
        return {"created": created, "updated": updated, "appended": 0, "error": True}


def _readable_headers():
    # One column per spec field so products can be added/edited by hand. The
    # full JSON stays last so markers like "_source" survive backup/restore.
    base = ["id", "model_no", "gem_product_id", "category", "image", "created_at", "updated_at",
            "description", "os", "processor", "ram", "storage"]
    return base + [f for f in _spec_fields() if f not in base] + ["extra_specs"]


def _write_backup(path, products, stamp):
    """Old sheet + every given product, saved as Desktop_Product.old-<stamp>.xlsx."""
    backup_path = os.path.join(os.path.dirname(path), f"Desktop_Product.old-{stamp}.xlsx")
    workbook, sheet, headers = _open(path)
    rows = _key_rows(sheet, headers)
    for product in products:
        for row in rows.get(_key_of(product)) or [sheet.max_row + 1]:
            _write_row(sheet, headers, row, _row_values(product))
    workbook.save(backup_path)
    return backup_path


def replace_directory(entries, dry_run=False):
    """Make the GeM products the whole Desktop directory and a fresh sheet.

    entries: [{"model_no", "gem_product_id", "specs", "summary"}], one per
    GeM listing (a model number can repeat). Every other Desktop product
    (bid-created models not yet on GeM included) is saved to a dated backup
    sheet and removed. Returns counts (and the backup file name)."""
    from .models import CatalogueProduct
    from .views.Desktop import _normalize_extra_specs

    path = excel_path()
    wanted = {product_key(e["model_no"], e.get("gem_product_id")): e for e in entries}
    with _file_lock(path):
        products = {_key_of(p): p for p in CatalogueProduct.objects.all()}
        matched = _match_products(wanted, products)
        clashes = sorted({k[0] for k, p in matched.items() if not is_desktop(p)})
        kept = {p.pk for p in matched.values()}
        removed = [p for p in products.values() if is_desktop(p) and p.pk not in kept]
        result = {
            "gem_products": sum(1 for k in wanted if k[0] not in clashes),
            "new": sum(1 for k in wanted if k not in matched),
            "updated": sum(1 for k in wanted if k in matched and k[0] not in clashes),
            "removed": len(removed),
            "skipped_non_desktop": clashes,
        }
        if dry_run:
            return result

        stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
        backup_path = _write_backup(path, [p for p in products.values() if is_desktop(p)], stamp)
        with paused(), transaction.atomic():
            for key, entry in wanted.items():
                if key[0] in clashes:
                    continue
                product = matched.get(key)
                if product is None:
                    product = CatalogueProduct(model_no=key[0])
                    old_specs = {}
                else:
                    old_specs = _product_specs(product)
                product.gem_product_id = key[1]
                # GeM has no "Motherboard / Chipset"; keep the directory's value.
                specs = _normalize_extra_specs({**old_specs, **entry["specs"]})
                if entry.get("sections"):
                    specs["_gem_sections"] = entry["sections"]
                product.extra_specs = specs
                product.category = "Desktop"
                for field, value in entry["summary"].items():
                    if value:
                        setattr(product, field, value)
                product.save()
            CatalogueProduct.objects.filter(pk__in=[p.pk for p in removed]).delete()

        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "Desktop Products"
        headers = _readable_headers()
        sheet.append(headers)
        keep = CatalogueProduct.objects.exclude(pk__in=[p.pk for p in removed]).order_by("model_no")
        for product in keep:
            if is_desktop(product):
                _write_row(sheet, headers, sheet.max_row + 1, _row_values(product))
        _save_workbook(workbook, path)
        with open(_baseline_marker(), "w", encoding="utf-8") as marker:
            marker.write(datetime.now().isoformat())
        _remember_mtime(path)
    result["backup_file"] = os.path.basename(backup_path)
    return result


def restore_from_backup(backup_path, model_nos=None):
    """Bring products back from a backup sheet (all, or only model_nos).
    Products already in the directory are left as they are."""
    from .models import CatalogueProduct

    workbook = openpyxl.load_workbook(backup_path, read_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    headers = [_cell_text(h) for h in next(rows)]
    only = {m.strip().upper() for m in model_nos} if model_nos else None
    existing = {product_key(m, g) for m, g in CatalogueProduct.objects.values_list("model_no", "gem_product_id")}
    restored = []
    for row in rows:
        data, specs = _row_product_data(headers, row)
        model_no = data.get("model_no", "").strip()
        key = product_key(model_no, data.get("gem_product_id"))
        if not key[0] or key in existing or (only and key[0] not in only):
            continue
        # Normal save: the signal writes each restored product into the sheet.
        with transaction.atomic():
            CatalogueProduct.objects.create(
                model_no=model_no, gem_product_id=key[1], category=data.get("category") or "Desktop",
                extra_specs=specs, image=data.get("image") or None,
                **{f: data.get(f, "") for f in TEXT_FIELDS if f != "category"},
            )
        existing.add(key)
        restored.append(model_no)
    workbook.close()
    return restored


def update_directory(entries, dry_run=False):
    """Apply GeM specs without removing anything: matching Desktop products get
    the GeM specs, unknown models are added. Excel is written once."""
    from .models import CatalogueProduct
    from .views.Desktop import _normalize_extra_specs

    path = excel_path()
    with _file_lock(path):
        products = {_key_of(p): p for p in CatalogueProduct.objects.all()}
        keys = [product_key(e["model_no"], e.get("gem_product_id")) for e in entries]
        matched = _match_products(keys, products)
        result = {"new": 0, "updated": 0, "unchanged": 0, "skipped_non_desktop": []}
        changed_products = []
        for key, entry in zip(keys, entries):
            product = matched.get(key)
            if product is not None and not is_desktop(product):
                result["skipped_non_desktop"].append(key[0])
                continue
            old_specs = _product_specs(product) if product else {}
            specs = _normalize_extra_specs({**old_specs, **entry["specs"]})
            if entry.get("sections"):
                specs["_gem_sections"] = entry["sections"]
            summary = {f: v for f, v in entry["summary"].items() if v}
            # Compare with the stored values as they are; a bid-created model
            # found on GeM also loses its "not yet on GeM" marker.
            same_specs = all((old_specs.get(k) or "") == (v or "") for k, v in specs.items())
            if product is not None and same_specs and "_source" not in old_specs and product.gem_product_id == key[1] and all(
                (getattr(product, f) or "") == v for f, v in summary.items()
            ):
                result["unchanged"] += 1
                continue
            result["updated" if product else "new"] += 1
            if dry_run:
                continue
            product = product or CatalogueProduct(model_no=key[0], category="Desktop")
            product.gem_product_id = key[1]
            product.extra_specs = specs
            for field, value in summary.items():
                setattr(product, field, value)
            changed_products.append(product)
        if dry_run or not changed_products:
            return result

        with paused(), transaction.atomic():
            for product in changed_products:
                product.save()
        workbook, sheet, headers = _open(path)
        rows = _key_rows(sheet, headers)
        for product in changed_products:
            # A legacy row (same model, no GeM id yet) becomes this listing's row.
            key = _key_of(product)
            row_numbers = rows.pop(key, None) or rows.pop((key[0], ""), None)
            for row in row_numbers or [sheet.max_row + 1]:
                _write_row(sheet, headers, row, _row_values(product))
        _save_workbook(workbook, path)
        _remember_mtime(path)
    return result
