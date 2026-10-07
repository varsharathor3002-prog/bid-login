import json
import os
import shutil
import tempfile
from unittest import mock

import openpyxl
from django.core import signing
from django.test import TestCase

from . import desktop_excel
from .models import CatalogueProduct, User

URL = "/api/catalogue/import-gem-specs/"


class GemProductSpecsImportTests(TestCase):
    def setUp(self):
        analyser = User.objects.create(
            username="specs-analyser", email="specs@example.com", password="password", role="analyser",
        )
        token = signing.dumps({"user_id": analyser.id, "role": analyser.role}, salt="gem-api-auth")
        self.auth = {"HTTP_AUTHORIZATION": f"Bearer {token}"}
        # Every save mirrors into Desktop_Product.xlsx; use a copy, never the real file.
        self.temp_dir = tempfile.mkdtemp()
        self.excel = os.path.join(self.temp_dir, "Desktop_Product.xlsx")
        workbook = openpyxl.Workbook()
        workbook.active.append(desktop_excel.DEFAULT_HEADERS)
        for i in range(1, 4):
            workbook.active.append([str(i), f"ACL-SHEET-{i}", "Desktop", None, "", "", json.dumps({"Type of RAM": "DDR4 RAM"}), "", "", "i3", "8 DDR4 RAM", "256"])
        workbook.save(self.excel)
        for target, value in [
            ("accounts.desktop_excel.excel_path", self.excel),
            ("accounts.desktop_excel.enabled", True),
        ]:
            patcher = mock.patch(target, return_value=value)
            patcher.start()
            self.addCleanup(patcher.stop)
        desktop_excel._last_seen_mtime = None
        # Start as an already-synced server; the first-sync test removes this.
        with open(f"{self.excel}.synced", "w") as marker:
            marker.write("test")
        self.addCleanup(shutil.rmtree, self.temp_dir, True)

    def excel_rows(self, model_no):
        sheet = openpyxl.load_workbook(self.excel).active
        headers = [cell.value for cell in sheet[1]]
        return [dict(zip(headers, row)) for row in sheet.iter_rows(min_row=2, values_only=True)
                if str(row[headers.index("model_no")] or "").upper() == model_no.upper()]

    def post(self, body):
        return self.client.post(URL, json.dumps(body), content_type="application/json", **self.auth)

    def test_requires_login(self):
        response = self.client.post(URL, "{}", content_type="application/json")
        self.assertEqual(response.status_code, 401)

    def test_updates_only_directory_fields_and_keeps_missing_ones(self):
        product = CatalogueProduct.objects.create(
            model_no="ACL-TEST-1", category="Desktop",
            extra_specs={"Type of RAM": "DDR4 RAM", "Number of HDMI Ports": "0", "Optical Drive": "No", "_source": "desktop_bid"},
        )
        pairs = [
            ["Model Number", "acl-test-1"],
            ["Number of HDMI Ports *", "1"],
            ["Type of RAM:", "DDR5"],
            ["Factory Pre-loaded Operating System by Desktop OEM", "Windows 11 Pro"],
            ["Size of Memory in Case of Dedicated Graphic Card (GB)", "4"],
            ["Seller Rating", "4.5"],
        ]
        preview = self.post({"pairs": pairs, "dry_run": True}).json()
        self.assertTrue(preview["exists"])
        self.assertFalse(preview["saved"])
        self.assertEqual(preview["found"], 4)
        product.refresh_from_db()
        self.assertEqual(product.extra_specs["Number of HDMI Ports"], "0")

        result = self.post({"pairs": pairs}).json()
        self.assertTrue(result["saved"])
        product.refresh_from_db()
        specs = product.extra_specs
        self.assertEqual(specs["Number of HDMI Ports"], "1")
        self.assertEqual(specs["Type of RAM"], "DDR5")
        self.assertEqual(specs["Factory Pre-loaded Operating System by DesktopOEM"], "Windows 11 Pro")
        self.assertEqual(specs["Size of Memory in Case of Dedicated Graphic Card(GB)"], "4")
        self.assertEqual(specs["Optical Drive"], "No")
        self.assertNotIn("Seller Rating", specs)
        self.assertNotIn("_source", specs)
        self.assertEqual(product.os, "Windows 11 Pro")

    def test_creates_missing_product_from_typed_model(self):
        result = self.post({"model_no": "acl-new-2", "pairs": [["Processor Number", "i5-14500"]]}).json()
        self.assertTrue(result["created"])
        product = CatalogueProduct.objects.get(model_no="ACL-NEW-2")
        self.assertEqual(product.category, "Desktop")
        self.assertEqual(product.processor, "i5-14500")

    def test_rejects_non_desktop_product_and_page_without_model(self):
        CatalogueProduct.objects.create(model_no="AIO-1", category="aio", extra_specs={"RAM": "8 GB"})
        response = self.post({"model_no": "AIO-1", "pairs": [["Type of RAM", "DDR5"]]})
        self.assertEqual(response.status_code, 400)
        self.assertEqual(CatalogueProduct.objects.get(model_no="AIO-1").extra_specs, {"RAM": "8 GB"})
        self.assertEqual(self.post({"pairs": [["Type of RAM", "DDR5"]]}).status_code, 400)

    def test_mirrors_existing_excel_row_without_adding_rows(self):
        sheet = openpyxl.load_workbook(self.excel).active
        total_rows = sheet.max_row
        model_no = sheet.cell(2, 2).value
        CatalogueProduct.objects.get_or_create(model_no=model_no, defaults={"category": "Desktop", "extra_specs": {}})
        result = self.post({"model_no": model_no, "pairs": [["Number of HDMI Ports", "7"]]}).json()
        self.assertTrue(result["excel_saved"], result.get("excel_error"))
        self.assertEqual(openpyxl.load_workbook(self.excel).active.max_row, total_rows)
        rows = self.excel_rows(model_no)
        self.assertEqual(len(rows), 1)
        self.assertEqual(json.loads(rows[0]["extra_specs"])["Number of HDMI Ports"], "7")

    def test_appends_new_product_to_excel(self):
        self.post({"model_no": "ACL-NEW-XL", "pairs": [["Processor Number", "i7-14700"]]})
        rows = self.excel_rows("ACL-NEW-XL")
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["processor"], "i7-14700")
        self.assertEqual(rows[0]["category"], "Desktop")

    def test_locked_excel_still_saves_directory(self):
        with mock.patch("accounts.desktop_excel.os.replace", side_effect=PermissionError):
            result = self.post({"model_no": "ACL-LOCK-1", "pairs": [["Type of RAM", "DDR5"]]}).json()
        self.assertTrue(result["saved"])
        self.assertFalse(result["excel_saved"])
        self.assertIn("open in another program", result["excel_error"])
        self.assertTrue(CatalogueProduct.objects.filter(model_no="ACL-LOCK-1").exists())
        self.assertEqual([f for f in os.listdir(self.temp_dir) if not f.startswith("Desktop_Product.xlsx")], [])


class DesktopExcelSyncTests(GemProductSpecsImportTests):
    """UI and bid-flow changes reach Desktop_Product.xlsx; Excel rows reach the directory."""

    def ui_create(self, model_no, **fields):
        with self.captureOnCommitCallbacks(execute=True):
            return self.client.post("/api/catalogue/create/", {"model_no": model_no, "category": "Desktop", **fields})

    def append_excel_row(self, values):
        workbook = openpyxl.load_workbook(self.excel)
        sheet = workbook.active
        headers = [cell.value for cell in sheet[1]]
        sheet.append([values.get(h) for h in headers])
        workbook.save(self.excel)
        future = os.path.getmtime(self.excel) + 5
        os.utime(self.excel, (future, future))

    def test_ui_create_rename_and_delete_follow_in_excel(self):
        self.assertEqual(self.ui_create("ACL-UI-1", processor="i3-12100").status_code, 201)
        self.assertEqual(self.excel_rows("ACL-UI-1")[0]["processor"], "i3-12100")

        product = CatalogueProduct.objects.get(model_no="ACL-UI-1")
        with self.captureOnCommitCallbacks(execute=True):
            self.client.post(f"/api/catalogue/{product.id}/update/", {"model_no": "ACL-UI-2", "ram": "16 DDR5"})
        self.assertEqual(self.excel_rows("ACL-UI-1"), [])
        self.assertEqual(self.excel_rows("ACL-UI-2")[0]["ram"], "16 DDR5")

        with self.captureOnCommitCallbacks(execute=True):
            self.client.delete(f"/api/catalogue/{product.id}/delete/")
        self.assertEqual(self.excel_rows("ACL-UI-2"), [])

    def test_row_added_in_excel_shows_in_directory(self):
        self.append_excel_row({
            "model_no": "ACL-XL-HAND", "category": "Desktop", "processor": "i5-13400",
            "extra_specs": json.dumps({"Type of RAM": "DDR5"}),
        })
        models = [p["model_no"] for p in self.client.get("/api/catalogue/?search=ACL-XL-HAND").json()]
        self.assertEqual(models, ["ACL-XL-HAND"])
        product = CatalogueProduct.objects.get(model_no="ACL-XL-HAND")
        self.assertEqual(product.processor, "i5-13400")
        self.assertEqual(desktop_excel._product_specs(product)["Type of RAM"], "DDR5")

    def test_row_edited_in_excel_updates_directory(self):
        self.ui_create("ACL-XL-EDIT", processor="old cpu")
        workbook = openpyxl.load_workbook(self.excel)
        sheet = workbook.active
        headers = [cell.value for cell in sheet[1]]
        for row in sheet.iter_rows(min_row=2):
            if row[headers.index("model_no")].value == "ACL-XL-EDIT":
                row[headers.index("processor")].value = "new cpu"
        workbook.save(self.excel)
        future = os.path.getmtime(self.excel) + 5
        os.utime(self.excel, (future, future))
        self.client.get("/api/catalogue/")
        self.assertEqual(CatalogueProduct.objects.get(model_no="ACL-XL-EDIT").processor, "new cpu")

    def test_directory_only_desktop_product_is_appended_to_excel(self):
        with desktop_excel.paused():
            CatalogueProduct.objects.create(model_no="ACL-DB-ONLY", category="Desktop", processor="i9")
            CatalogueProduct.objects.create(model_no="AIO-DB-ONLY", category="aio")
        self.client.get("/api/catalogue/")
        self.assertEqual(self.excel_rows("ACL-DB-ONLY")[0]["processor"], "i9")
        self.assertEqual(self.excel_rows("AIO-DB-ONLY"), [])

    def test_first_sync_keeps_directory_values_and_backs_up_excel(self):
        os.remove(f"{self.excel}.synced")
        model_no = openpyxl.load_workbook(self.excel).active.cell(2, 2).value
        with desktop_excel.paused():
            product, _ = CatalogueProduct.objects.get_or_create(model_no=model_no, defaults={"category": "Desktop"})
            product.processor = "DIRECTORY VALUE"
            product.save()
        future = os.path.getmtime(self.excel) + 5
        os.utime(self.excel, (future, future))
        self.client.get("/api/catalogue/")
        self.assertEqual(CatalogueProduct.objects.get(pk=product.pk).processor, "DIRECTORY VALUE")
        self.assertEqual(self.excel_rows(model_no)[0]["processor"], "DIRECTORY VALUE")
        self.assertTrue(os.path.exists(f"{self.excel}.synced"))
        self.assertTrue([f for f in os.listdir(self.temp_dir) if f.startswith("Desktop_Product.backup-")])


class GemMarketReplaceTests(GemProductSpecsImportTests):
    URL = "/api/catalogue/replace-from-gem-market/"

    def gem_product(self, model_no, ram="DDR5", brand="acxxel", listing=None):
        listing = listing or abs(hash(model_no)) % 10**9
        return {"url": f"https://mkp.gem.gov.in/x/{model_no.lower()}/p-5116877-{listing}-cat.html#variant",
                "brand": brand, "model_no": model_no,
                "pairs": [["Processor Number", "Intel Core i5- 14500"], ["Type of RAM", ram]]}

    def replace(self, products, **extra):
        body = {"products": products, "total": len(products), **extra}
        return self.client.post(self.URL, json.dumps(body), content_type="application/json", **self.auth)

    def setUp(self):
        super().setUp()
        with desktop_excel.paused():
            CatalogueProduct.objects.all().delete()
            CatalogueProduct.objects.create(model_no="OLD-1", category="Desktop", processor="old")
            CatalogueProduct.objects.create(model_no="GEM-1", category="Desktop", extra_specs={"Motherboard / Chipset": "H610"})
            CatalogueProduct.objects.create(model_no="BID-NEW", category="Desktop", extra_specs={"_source": "desktop_bid"})
            CatalogueProduct.objects.create(model_no="AIO-KEEP", category="aio", extra_specs={"RAM": "8"})

    def test_rejects_incomplete_scan(self):
        body = {"products": [self.gem_product("GEM-1")], "total": 142}
        response = self.client.post(self.URL, json.dumps(body), content_type="application/json", **self.auth)
        self.assertEqual(response.status_code, 400)
        self.assertTrue(CatalogueProduct.objects.filter(model_no="OLD-1").exists())

    def test_preview_changes_nothing(self):
        result = self.replace([self.gem_product("GEM-1"), self.gem_product("GEM-2")], dry_run=True).json()
        self.assertEqual((result["new"], result["updated"], result["removed"]), (1, 1, 2))
        self.assertTrue(CatalogueProduct.objects.filter(model_no="OLD-1").exists())
        self.assertFalse(CatalogueProduct.objects.filter(model_no="GEM-2").exists())

    def test_replace_backs_up_old_products_and_writes_fresh_excel(self):
        products = [self.gem_product("GEM-1", listing=1), self.gem_product("GEM-2", listing=2),
                    self.gem_product("GEM-2", ram="DDR4", listing=3)]
        result = self.replace(products).json()
        self.assertEqual(result["duplicates"], ["GEM-2"])
        models = sorted(CatalogueProduct.objects.values_list("model_no", flat=True))
        self.assertEqual(models, ["AIO-KEEP", "GEM-1", "GEM-2", "GEM-2"])
        # The existing GEM-1 is taken over by its GeM listing (keeps its chipset).
        gem1 = CatalogueProduct.objects.get(model_no="GEM-1")
        self.assertEqual(gem1.gem_product_id, "5116877-1")
        self.assertEqual(gem1.extra_specs["Type of RAM"], "DDR5")
        self.assertEqual(gem1.extra_specs["Motherboard / Chipset"], "H610")
        configs = dict(CatalogueProduct.objects.filter(model_no="GEM-2").values_list("gem_product_id", "extra_specs"))
        self.assertEqual({k: v["Type of RAM"] for k, v in configs.items()}, {"5116877-2": "DDR5", "5116877-3": "DDR4 RAM"})

        sheet = openpyxl.load_workbook(self.excel).active
        headers = [c.value for c in sheet[1]]
        self.assertIn("Type of RAM", headers)
        excel_models = sorted(r[headers.index("model_no")] for r in sheet.iter_rows(min_row=2, values_only=True))
        self.assertEqual(excel_models, ["GEM-1", "GEM-2", "GEM-2"])

        backup = os.path.join(self.temp_dir, result["backup_file"])
        self.assertTrue(os.path.exists(backup))
        restored = desktop_excel.restore_from_backup(backup, ["OLD-1", "BID-NEW"])
        self.assertEqual(sorted(restored), ["BID-NEW", "OLD-1"])
        self.assertEqual(CatalogueProduct.objects.get(model_no="BID-NEW").extra_specs.get("_source"), "desktop_bid")
        self.assertEqual(CatalogueProduct.objects.get(model_no="OLD-1").processor, "old")

    def test_rejects_other_brand(self):
        response = self.replace([self.gem_product("GEM-1", brand="acer")])
        self.assertEqual(response.status_code, 400)

    def test_update_mode_changes_specs_and_removes_nothing(self):
        products = [self.gem_product("GEM-1"), self.gem_product("GEM-NEW")]
        preview = self.replace(products, mode="update", dry_run=True).json()
        self.assertEqual((preview["updated"], preview["new"], preview["unchanged"]), (1, 1, 0))
        self.assertFalse(CatalogueProduct.objects.filter(model_no="GEM-NEW").exists())

        result = self.replace(products, mode="update").json()
        self.assertEqual(result["mode"], "update")
        models = set(CatalogueProduct.objects.values_list("model_no", flat=True))
        self.assertEqual(models, {"OLD-1", "GEM-1", "GEM-NEW", "BID-NEW", "AIO-KEEP"})
        gem1 = CatalogueProduct.objects.get(model_no="GEM-1")
        self.assertEqual(gem1.extra_specs["Type of RAM"], "DDR5")
        self.assertEqual(gem1.extra_specs["Motherboard / Chipset"], "H610")
        self.assertEqual(self.excel_rows("GEM-NEW")[0]["processor"], "Intel Core i5- 14500")
        self.assertEqual(json.loads(self.excel_rows("GEM-1")[0]["extra_specs"])["Type of RAM"], "DDR5")

        again = self.replace(products, mode="update", dry_run=True).json()
        self.assertEqual((again["updated"], again["new"], again["unchanged"]), (0, 0, 2))

    def test_update_mode_accepts_a_single_model(self):
        body = {"products": [self.gem_product("GEM-1")], "total": 142, "mode": "update"}
        response = self.client.post(self.URL, json.dumps(body), content_type="application/json", **self.auth)
        self.assertEqual(response.status_code, 200)


class HighEndMappingTests(TestCase):
    def test_high_end_labels_map_to_directory_fields(self):
        from .views.GemProductSpecs import _map_pairs
        specs, _ = _map_pairs([
            ["Base Processor Number", "NA for Higher Processor"],
            ["Higher Processor Number", "Intel Core i7- 13700"],
            ["Out of Band Management", "NA"],
            ["Factory Pre-loaded Operating System", "Window 11 Professional"],
            ["Type of Storage Installed with the System", "NVME - SSD"],
            ["Primary Storage Capacity (in GB)", "1024"],
            ["Availability of Secondary Storage", "HDD@7200RPM"],
            ["Secondary Storage Capacity (in GB)", "2048"],
            ["Number of DIMM Slots populated with Memory Card", "1"],
            ["Number of USB Type A Ports (Version 2 Point 0)", "4"],
        ])
        self.assertEqual(specs["Processor Number"], "Intel Core i7- 13700")
        self.assertEqual(specs["Computer Type"], "High End")
        self.assertEqual(specs["Factory Pre-loaded Operating System by DesktopOEM"], "Window 11 Professional")
        self.assertEqual(specs["SSD - Storage Capacity (in GB)"], "1024")
        self.assertEqual(specs["HDD - Storage Capacity (in GB)"], "2048")
        self.assertEqual(specs["Number of DIMM Slots Populated with MemoryCard/Module"], "1")
        self.assertEqual(specs["Number of USB Type A Port (Version 2 Point 0)"], "4")
        self.assertNotIn("Out of Band", str(specs))

    def test_no_secondary_storage_means_zero_hdd(self):
        from .views.GemProductSpecs import _map_pairs
        specs, _ = _map_pairs([
            ["Base Processor Number", "Intel Core i5- 14500"],
            ["Type of Storage Installed with the System", "NVME - SSD"],
            ["Primary Storage Capacity (in GB)", "512"],
            ["Availability of Secondary Storage", "No Secondary Storage"],
            ["Secondary Storage Capacity (in GB)", "0"],
        ])
        self.assertEqual(specs["Processor Number"], "Intel Core i5- 14500")
        self.assertEqual(specs["HDD - Storage Capacity (in GB)"], "0")


class GemCategoryExcelTests(GemProductSpecsImportTests):
    URL = "/api/catalogue/gem-category-excel/"

    def save(self, products, category="All in One PC (V2)"):
        body = {"category": category, "products": products}
        return self.client.post(self.URL, json.dumps(body), content_type="application/json", **self.auth)

    def test_writes_one_excel_per_category_and_updates_by_model(self):
        target = os.path.join(self.temp_dir, "All_in_One_PC_V2.xlsx")
        with mock.patch("accounts.views.GemProductSpecs.category_excel_path", return_value=target):
            first = self.save([
                {"model_no": "BQ2700", "brand": "acxxel", "title": "acxxel AIO", "url": "https://mkp.gem.gov.in/a/p-1-2-cat.html#v",
                 "pairs": [["Processor Number", "Intel Core i3-12100"], ["Screen Size", "23.8"]]},
                {"model_no": "", "brand": "acxxel", "title": "No model", "url": "https://mkp.gem.gov.in/b/p-3-4-cat.html",
                 "pairs": [["Processor Number", "i5"]]},
                {"model_no": "HP-1", "brand": "HP", "title": "HP AIO", "url": "https://mkp.gem.gov.in/c/p-5-6-cat.html",
                 "pairs": [["Processor Number", "i7"]]},
            ]).json()
            self.assertEqual((first["added"], first["updated"]), (2, 0))
            second = self.save([
                {"model_no": "bq2700", "brand": "acxxel", "title": "acxxel AIO", "url": "https://mkp.gem.gov.in/a/p-1-2-cat.html",
                 "pairs": [["Processor Number", "Intel Core i5-12400"], ["Webcam", "Yes"]]},
            ]).json()
            self.assertEqual((second["added"], second["updated"], second["total_rows"]), (0, 1, 2))
        sheet = openpyxl.load_workbook(target).active
        headers = [c.value for c in sheet[1]]
        self.assertEqual(headers[:5], ["model_no", "brand", "product_name", "url", "updated_at"])
        self.assertIn("Webcam", headers)
        row = [r for r in sheet.iter_rows(min_row=2, values_only=True) if r[0] == "bq2700"][0]
        self.assertEqual(row[headers.index("Processor Number")], "Intel Core i5-12400")
        self.assertEqual(row[headers.index("Screen Size")], "23.8")
        self.assertEqual(row[headers.index("url")], "https://mkp.gem.gov.in/a/p-1-2-cat.html")
        self.assertEqual(sheet.max_row, 3)

    def test_category_file_names_are_safe(self):
        from .views.GemProductSpecs import category_excel_path
        path = category_excel_path("Toner Cartridges / Ink Cartridges / Consumables for Printers")
        self.assertTrue(path.endswith(os.path.join("gem_products", "Toner_Cartridges_Ink_Cartridges_Consumables_for_Printers.xlsx")))

    def test_other_brand_rows_are_removed_and_rejected(self):
        target = os.path.join(self.temp_dir, "Toner.xlsx")
        workbook = openpyxl.Workbook()
        workbook.active.append(["model_no", "brand", "product_name", "url", "updated_at"])
        workbook.active.append(["OLD-HP", "HP", "", "https://x/1", ""])
        workbook.active.append(["OLD-ACX", "acxxel", "", "https://x/2", ""])
        workbook.save(target)
        with mock.patch("accounts.views.GemProductSpecs.category_excel_path", return_value=target):
            only_others = self.save([{"model_no": "HP-2", "brand": "HP", "url": "https://x/3", "pairs": []}], category="Toner")
            self.assertEqual(only_others.status_code, 400)
            result = self.save([{"model_no": "ACX-2", "brand": "acxxel", "url": "https://x/4", "pairs": [["Colour", "Black"]]}], category="Toner").json()
        self.assertEqual(result["removed_other_brands"], 1)
        models = [r[0] for r in openpyxl.load_workbook(target).active.iter_rows(min_row=2, values_only=True)]
        self.assertEqual(sorted(models), ["ACX-2", "OLD-ACX"])

    def test_aio_products_also_go_to_directory_aio_tab(self):
        target = os.path.join(self.temp_dir, "All_in_One_PC_V2.xlsx")
        pairs = [
            ["Processor Number", "Intel Core i5-12400"],
            ["Operating System (Factory Preloaded with Certification)", "Windows 11 Home"],
            ["Type of RAM", "DDR4"], ["RAM Size (GB)", "16"],
            ["Type of Storage Installed with the System", "NVMe SSD"], ["Storage Capacity (in GB)", "512"],
            ["Display Size - Diagonal (in Inches)", '58.1 - 63 (22.87" - 24.8")'],
            ["Type of In-built Wireless Connectivity", "Wi-Fi 5 (802.11ac) + Bluetooth 4.2"],
            ["Keyboard Connectivity", "Wired"],
            ["Number of Ports", "1-HDMI,1-DP"],
            ["BIS Registration Number", "R-93018430"],
        ]
        CatalogueProduct.objects.create(model_no="AXL-AIO-OLD", category="aio", extra_specs={"_source": "aio_bid", "RAM": "8GB"})
        with mock.patch("accounts.views.GemProductSpecs.category_excel_path", return_value=target):
            result = self.save([
                {"model_no": "axl-aio-test-1818", "brand": "acxxel", "title": "acxxel i5 AIO", "url": "https://x/1", "pairs": pairs},
                {"model_no": "AXL-AIO-OLD", "brand": "acxxel", "title": "old", "url": "https://x/2", "pairs": pairs},
            ]).json()
        self.assertEqual((result["directory_added"], result["directory_updated"]), (1, 1))
        product = CatalogueProduct.objects.get(model_no="AXL-AIO-TEST-1818")
        self.assertEqual(product.category, "aio")
        self.assertEqual((product.processor, product.ram, product.storage, product.os),
                         ("Intel Core i5 12400", "16GB DDR4", "512 GB NVMe SSD", "Windows 11 Home"))
        specs = product.extra_specs
        self.assertEqual(specs["Screen Size"], "22.87 to 24.8 inch")
        self.assertEqual(specs["WiFi Bluetooth"], "Wi-Fi 5 (802.11ac) + Bluetooth 4.2")
        self.assertEqual(specs["Keyboard Mouse"], "Wired")
        self.assertEqual(specs["Motherboard Ports"], "1-HDMI,1-DP")
        self.assertNotIn("_source", CatalogueProduct.objects.get(model_no="AXL-AIO-OLD").extra_specs)
        self.assertNotIn("BIS Registration Number", specs)
        listed = [p["model_no"] for p in self.client.get("/api/aio-catalogue/").json()]
        self.assertIn("AXL-AIO-TEST-1818", listed)


class GemConfigurationTests(GemMarketReplaceTests):
    """One model number, several GeM configurations: all are kept and found."""

    def test_update_keeps_every_configuration_and_search_finds_all(self):
        products = [self.gem_product("ACL-CFG", ram=r, listing=i) for i, r in [(11, "DDR4"), (12, "DDR5"), (13, "DDR4")]]
        result = self.replace(products, mode="update").json()
        self.assertEqual((result["new"], result["duplicates"]), (3, ["ACL-CFG"]))
        found = self.client.get("/api/catalogue/?search=ACL-CFG").json()
        self.assertEqual(sorted(p["gem_product_id"] for p in found), ["5116877-11", "5116877-12", "5116877-13"])
        self.assertEqual({p["model_no"] for p in found}, {"ACL-CFG"})
        rows = self.excel_rows("ACL-CFG")
        self.assertEqual(sorted(r["gem_product_id"] for r in rows), ["5116877-11", "5116877-12", "5116877-13"])

        again = self.replace(products, mode="update", dry_run=True).json()
        self.assertEqual((again["new"], again["updated"], again["unchanged"]), (0, 0, 3))

        # Deleting one configuration removes only its Excel row.
        one = CatalogueProduct.objects.get(model_no="ACL-CFG", gem_product_id="5116877-12")
        with self.captureOnCommitCallbacks(execute=True):
            self.client.delete(f"/api/catalogue/{one.id}/delete/")
        self.assertEqual(sorted(r["gem_product_id"] for r in self.excel_rows("ACL-CFG")), ["5116877-11", "5116877-13"])

    def test_aio_configurations_are_all_in_the_aio_tab(self):
        target = os.path.join(self.temp_dir, "All_in_One_PC_V2.xlsx")
        items = [
            {"model_no": "AXL-AIO-CFG", "brand": "acxxel", "title": f"AIO {cpu}",
             "url": f"https://mkp.gem.gov.in/a/p-5116877-{n}-cat.html", "pairs": [["Processor Number", cpu]]}
            for n, cpu in [(21, "Intel Core i5-12400"), (22, "Intel Core i7-14700")]
        ]
        with mock.patch("accounts.views.GemProductSpecs.category_excel_path", return_value=target):
            body = {"category": "All in One PC (V2)", "products": items}
            result = self.client.post("/api/catalogue/gem-category-excel/", json.dumps(body),
                                      content_type="application/json", **self.auth).json()
        self.assertEqual((result["directory_added"], result["added"]), (2, 2))
        listed = [p for p in self.client.get("/api/aio-catalogue/?search=axl-aio-cfg").json()]
        self.assertEqual(sorted(p["processor"] for p in listed), ["Intel Core i5 12400", "Intel Core i7 14700"])

    def test_toner_category_with_slashes_saves(self):
        target = os.path.join(self.temp_dir, "Toner.xlsx")
        category = "Toner Cartridges / Ink Cartridges / Consumables for Printers"
        with mock.patch("accounts.views.GemProductSpecs.category_excel_path", return_value=target):
            body = {"category": category, "products": [{"model_no": "ACX-T1", "brand": "acxxel",
                    "url": "https://mkp.gem.gov.in/t/p-1-9-cat.html", "pairs": [["Product Class of Cartridge", "Compatible"]]}]}
            response = self.client.post("/api/catalogue/gem-category-excel/", json.dumps(body),
                                        content_type="application/json", **self.auth)
        self.assertEqual(response.status_code, 200, response.content)
        self.assertEqual(openpyxl.load_workbook(target).active.title, "Toner Cartridges Ink Cartridges")

    def test_toner_products_go_to_the_toner_tab(self):
        target = os.path.join(self.temp_dir, "Toner.xlsx")
        pairs = [
            ["Product Class of Cartridge", "Compatible"],
            ["Printer/Multifunction Machines Brand for which offered Cartridge/Consumable is Suitable", "Canon"],
            ["Type of Cartridge/Consumable", "Toner Cartridge"],
            ["Color of the Ink/Toner", "Black"],
            ["Model Number of OEM's Printer/OEM's Multi Function Machine", "MF232w / MF237w"],
            ["Method for the Determination of Toner Cartridge Yield for Monochromatic Electrophotographic Printers and Multi-Function Devices that Contain Printer Components", "As per IS/ISO/IEC 19752 : 2017"],
            ["Minimum Yield of the Replacement Cartridge/Consumable offered (Number of Pages)", "1001 to 2000"],
        ]
        body = {"category": "Toner Cartridges / Ink Cartridges / Consumables for Printers", "products": [
            {"model_no": "AXL-337", "brand": "acxxel", "title": "acxxel Compatible Toner Cartridge",
             "url": "https://mkp.gem.gov.in/t/p-5116877-31-cat.html", "pairs": pairs}]}
        with mock.patch("accounts.views.GemProductSpecs.category_excel_path", return_value=target):
            result = self.client.post("/api/catalogue/gem-category-excel/", json.dumps(body),
                                      content_type="application/json", **self.auth).json()
        self.assertEqual(result["directory_added"], 1)
        listed = self.client.get("/api/toner-catalogue/?search=axl-337").json()
        self.assertEqual(len(listed), 1)
        toner = listed[0]
        self.assertEqual((toner["brand"], toner["cartridge_type"], toner["product_class"], toner["colour"],
                          toner["technology"], toner["page_yield"], toner["yield_standard"], toner["compatibility"]),
                         ("Canon", "Toner Cartridge", "Compatible", "Black", "Laser", "1001 to 2000", "ISO/IEC 19752", "MF232w / MF237w"))

    def test_each_product_keeps_its_own_gem_sections_without_certification(self):
        sections = [["GENERIC", [["Product Class of Cartridge", "Compatible"], ["Color of the Ink/Toner", "Black"]]],
                    ["Certification", [["ROHS Compliance", "Yes"]]]]
        toner = {"model_no": "AXL-SEC", "brand": "acxxel", "url": "https://mkp.gem.gov.in/t/p-5116877-41-cat.html",
                 "pairs": sections[0][1] + sections[1][1], "sections": sections}
        with mock.patch("accounts.views.GemProductSpecs.category_excel_path", return_value=os.path.join(self.temp_dir, "T.xlsx")):
            body = {"category": "Toner Cartridges / Ink Cartridges / Consumables for Printers", "products": [toner]}
            self.client.post("/api/catalogue/gem-category-excel/", json.dumps(body), content_type="application/json", **self.auth)
        listed = self.client.get("/api/toner-catalogue/?search=axl-sec").json()[0]
        self.assertEqual(listed["extra_specs"]["_gem_sections"], [sections[0]])

        desktop = dict(self.gem_product("ACL-SEC", listing=42),
                       sections=[["PROCESSOR", [["Processor Number", "Intel Core i5- 14500"]]], ["CERTIFICATION", [["X", "Y"]]]])
        self.replace([desktop], mode="update")
        found = self.client.get("/api/catalogue/?search=ACL-SEC").json()[0]
        self.assertEqual(found["gem_sections"], [["PROCESSOR", [["Processor Number", "Intel Core i5- 14500"]]]])
